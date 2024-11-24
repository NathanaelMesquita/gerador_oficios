import os
from datetime import datetime
from docx import Document
from docx2pdf import convert
import streamlit as st

# Caminho do modelo DOCX
TEMPLATE_PATH = "modelo.docx"

# Função para preencher o modelo DOCX
def preencher_modelo(template_path, context, output_path):
    """
    Preenche o modelo DOCX substituindo os placeholders pelos valores fornecidos.
    """
    # Abrir o modelo
    doc = Document(template_path)
    
    # Substituir placeholders
    for paragraph in doc.paragraphs:
        for key, value in context.items():
            if f"{{{key}}}" in paragraph.text:
                paragraph.text = paragraph.text.replace(f"{{{key}}}", value)

    # Salvar o documento preenchido
    doc.save(output_path)

def gerar_pdf(docx_path, pdf_path):
    """
    Converte o DOCX preenchido para PDF utilizando docx2pdf.
    """
    convert(docx_path, pdf_path)

def gerar_oficio(context):
    """
    Preenche o modelo DOCX, converte para PDF e oferece download.
    """
    # Caminho para salvar o DOCX preenchido
    preenchido_path = "oficio_preenchido.docx"
    
    # Caminho para salvar o PDF
    pdf_path = "oficio_preenchido.pdf"

    # Preencher o modelo
    preencher_modelo(TEMPLATE_PATH, context, preenchido_path)

    # Converter para PDF
    gerar_pdf(preenchido_path, pdf_path)

    # Oferecer o PDF para download
    with open(pdf_path, "rb") as pdf_file:
        st.download_button(
            label="Baixar Ofício em PDF",
            data=pdf_file,
            file_name="oficio_preenchido.pdf",
            mime="application/pdf",
        )

# Interface Streamlit
st.title("Gerador de Ofícios Automáticos")
st.header("Preencha os dados abaixo para gerar o ofício:")

# Entradas do usuário
numero_oficio = st.text_input("Número do Ofício", "001")
ano_atual = datetime.now().year
operadora = st.selectbox("Operadora", ["Vivo", "Claro", "TIM"])
tipo_referencia = st.selectbox("Tipo de Referência", ["IP", "VPI"])
numero_referencia = st.text_input(f"Número do {tipo_referencia}", "12345")
ano_referencia = st.text_input("Ano do IP/VPI", str(ano_atual))
cpf = st.text_area("Lista de CPFs (um por linha)", "000.000.000-00")
email = st.text_input("E-mail para Resposta", "delegado@exemplo.com")
nome_delegado = st.text_input("Nome do Delegado", "Delegado Exemplo")

# Gerar Ofício
if st.button("Gerar Ofício"):
    # Criar o contexto para substituição
    context = {
        "número_oficio": numero_oficio,
        "ano_atual": str(ano_atual),
        "operadora": operadora,
        "IP/VPI": tipo_referencia,
        "Número_ip_vpi": numero_referencia,
        "ano_ip_vpi": ano_referencia,
        "email_delegacia": email,
        "nome_delegado": nome_delegado,
        "cpf_lista": "\n".join([f"  - {c.strip()}" for c in cpf.splitlines() if c.strip()]),
    }

    # Gerar o documento e permitir download
    gerar_oficio(context)
